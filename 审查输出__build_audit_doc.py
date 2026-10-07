from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "审查输出"
EVIDENCE = OUT_DIR / "evidence"
OUTPUT = OUT_DIR / "严格审稿与真实性证据报告.docx"

NAVY = "17324D"
TEAL = "1D6F78"
BLUE = "2A5D85"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E8F4F3"
PALE_RED = "FCE8E6"
PALE_ORANGE = "FFF1DB"
PALE_YELLOW = "FFF8D9"
PALE_GREEN = "EAF5EA"
PALE_GRAY = "F2F4F5"
MID_GRAY = "667581"
DARK = "1F2933"
WHITE = "FFFFFF"
RED = "A83232"
ORANGE = "B56212"
GREEN = "2E6B3C"


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


AUDIT = load_json(EVIDENCE / "audit_evidence.json")
REBUILD = load_json(EVIDENCE / "raw_rebuild_evidence.json")
SOURCES = load_json(EVIDENCE / "source_endpoint_audit.json")
FIG_REPRO = load_json(OUT_DIR / "figure_reproduction" / "figure_reproduction_evidence.json")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm: float):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, name="Microsoft YaHei", size=9.5, bold=False, color=DARK, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_para_format(paragraph, after=4, before=0, line=1.2, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    fmt.widow_control = True


def text_para(doc, text, *, style=None, bold_prefix=None, after=5, before=0, line=1.25, color=DARK, size=9.5, keep=False):
    p = doc.add_paragraph(style=style)
    set_para_format(p, after=after, before=before, line=line, keep=keep)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True, color=color, size=size)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, color=color, size=size)
    else:
        r = p.add_run(text)
        set_run_font(r, color=color, size=size)
    return p


def bullet(doc, text, level=0, *, color=DARK):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_para_format(p, after=3, line=1.18)
    r = p.add_run(text)
    set_run_font(r, color=color, size=9.3)
    return p


def numbered(doc, text, level=0):
    p = doc.add_paragraph(style="List Number" if level == 0 else "List Number 2")
    set_para_format(p, after=3, line=1.18)
    r = p.add_run(text)
    set_run_font(r, size=9.3)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    return p


def add_callout(doc, title, text, *, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(17.2)
    cell = table.cell(0, 0)
    set_cell_width(cell, 17.2)
    set_cell_margins(cell, top=130, start=180, bottom=130, end=180)
    shade_cell(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    set_para_format(p, after=2, line=1.15)
    r = p.add_run(title)
    set_run_font(r, bold=True, color=accent, size=10)
    p2 = cell.add_paragraph()
    set_para_format(p2, after=0, line=1.22)
    r2 = p2.add_run(text)
    set_run_font(r2, color=DARK, size=9.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_width(cell, 17.0)
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    shade_cell(cell, "F5F6F7")
    p = cell.paragraphs[0]
    set_para_format(p, after=0, line=1.05)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=7.8, color="28323A")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None, *, header_fill=NAVY, font_size=8.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, h in enumerate(headers):
        cell = hdr.cells[j]
        shade_cell(cell, header_fill)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        if widths:
            set_cell_width(cell, widths[j])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para_format(p, after=0, line=1.0)
        r = p.add_run(str(h))
        set_run_font(r, bold=True, color=WHITE, size=8.2)
    for i, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for j, value in enumerate(row_data):
            cell = row.cells[j]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[j])
            if i % 2:
                shade_cell(cell, "F7F9FA")
            p = cell.paragraphs[0]
            set_para_format(p, after=0, line=1.12)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_status_table(doc, rows):
    table = add_table(doc, ["审查维度", "状态", "证据与判定"], rows, [4.0, 2.3, 10.7], font_size=8.5)
    color_map = {"通过": PALE_GREEN, "部分通过": PALE_YELLOW, "不通过": PALE_RED, "待作者确认": PALE_ORANGE}
    for row in table.rows[1:]:
        status = row.cells[1].text.strip()
        shade_cell(row.cells[1], color_map.get(status, PALE_GRAY))
        for run in row.cells[1].paragraphs[0].runs:
            set_run_font(run, bold=True, color=GREEN if status == "通过" else RED if status == "不通过" else ORANGE, size=8.5)
    return table


def add_figure(doc, image_path: Path, caption: str, width_cm=16.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_format(c, after=7, line=1.12)
    r = c.add_run(caption)
    set_run_font(r, size=8.2, italic=True, color=MID_GRAY)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("审查文档  |  ")
    set_run_font(r, size=8, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r._r.append(fld_char1)
    r._r.append(instr)
    r._r.append(fld_char2)


def configure_document(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.65)
    sec.bottom_margin = Cm(1.55)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.22

    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 10, 6),
        ("Heading 2", 12, TEAL, 8, 4),
        ("Heading 3", 10.5, BLUE, 6, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        st = doc.styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(9.3)

    # Running header and footer.
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run("PLUVIAL FLOOD RISK · INDEPENDENT TECHNICAL AUDIT")
    set_run_font(r, name="Arial", size=7.5, bold=True, color=TEAL)
    p_pr = hp._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), TEAL)
    pbdr.append(bottom)
    p_pr.append(pbdr)
    add_page_number(sec.footer.paragraphs[0])


def cover(doc: Document):
    # Memo masthead: restrained metadata + strong rule.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("TECHNICAL AUDIT MEMORANDUM")
    set_run_font(r, name="Arial", size=8.5, bold=True, color=TEAL)

    p = doc.add_paragraph()
    set_para_format(p, after=5, line=1.0)
    r = p.add_run("严格审稿与真实性证据报告")
    set_run_font(r, size=24, bold=True, color=NAVY)

    p = doc.add_paragraph()
    set_para_format(p, after=14, line=1.2)
    r = p.add_run("代码、原始/处理数据、图表—结果逻辑与论文表述的交叉核验")
    set_run_font(r, size=12, color=TEAL)

    meta = doc.add_table(rows=5, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    items = [
        ("项目", "20260522-pluvial-flood-risk-DGGS-H3"),
        ("审查对象", "docs/paper/manuscript.pdf、report.pdf、源代码、raw/processed 数据、模型与输出图表"),
        ("审查结论", "重大修改（Major Revision）— 当前版本不建议投稿"),
        ("证据快照", f"Git HEAD {AUDIT['git']['head'][:12]}；工作树 dirty={AUDIT['git']['dirty']}"),
        ("报告生成", datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M %Z")),
    ]
    for i, (k, v) in enumerate(items):
        set_cell_width(meta.cell(i, 0), 3.2)
        set_cell_width(meta.cell(i, 1), 13.6)
        set_cell_margins(meta.cell(i, 0), top=90, bottom=90)
        set_cell_margins(meta.cell(i, 1), top=90, bottom=90)
        shade_cell(meta.cell(i, 0), NAVY)
        p0 = meta.cell(i, 0).paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        rr = p0.add_run(k)
        set_run_font(rr, bold=True, color=WHITE, size=8.8)
        p1 = meta.cell(i, 1).paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        rr = p1.add_run(v)
        set_run_font(rr, bold=(k == "审查结论"), color=RED if k == "审查结论" else DARK, size=8.8)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_callout(
        doc,
        "一句话结论",
        "现有证据支持“当前数值是由本项目代码针对本地原始快照计算得到，未发现引用论文 PFIb 数据或实质段落被直接移入”的有限结论；但无法支持“所有输入均为作者自有数据”“数据端点均为官方来源”“图 2(c) 是全量拟合”“当前提交版本已被不可变归档”等更强声明。必须先修复 P0 项并整链重跑。",
        fill=PALE_RED,
        accent=RED,
    )

    text_para(
        doc,
        "证据边界：本报告是对当前工作目录快照的技术审查，不是数据权属法律意见，也不能对所有外部文献做全宇宙查重。公开政府/第三方数据本来就不属于作者；应主张“分析、代码执行与派生结果由作者完成”，而不是“原始数据归作者所有”。",
        color=MID_GRAY,
        size=8.7,
        after=3,
    )

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def section_break(doc):
    # Let major sections flow naturally. Heading styles already keep the heading
    # with the following paragraph; forced page breaks created sparse orphan pages.
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(2)


def build_document():
    doc = Document()
    configure_document(doc)
    cover(doc)

    add_heading(doc, "1. 总体审稿结论", 1)
    add_status_table(
        doc,
        [
            ("当前表格→指标独立复算", "通过", "141/956 个网格的空间 CV、OOF ROC-AUC/AP、准确率、F1、R²均可从现有表格独立复算并与 metadata 一致。"),
            ("当前 raw→processed 重建", "通过", "关闭 synthetic fallback 后，H3 集合、标签、文本字段一致；连续几何量最大差异约 3.7×10⁻⁸，属于浮点几何差异。"),
            ("参考论文数据未被直接使用", "部分通过", "未发现 PFIb 数据/模型实现或可疑实质文字复用；但只检查了仓库内提供的参考论文，不能证明对所有来源的绝对原创。"),
            ("来源为官方且权属可追踪", "不通过", "DEP 和 311 当前 ArcGIS 服务由个人/非机构账号持有，缺许可证与官方组织标识；manifest 缺精确 URL、逐层时间、许可和 SHA-256。"),
            ("论文边界框与数据一致", "不通过", "论文写的 bbox 生成 262 个 R9 网格，而 141 个结果实际来自较小 smoke bbox。"),
            ("图 2(c)/adaptive 为全量拟合", "不通过", "保存模型与 80% split 重拟合逐点完全相同，与 100% 全量拟合显著不同。"),
            ("图片→数据→代码前后对应", "部分通过", "图可由当前工作流再生成，核心 CSV 数值吻合；但 Fig.2(c) 语义错误，Fig.4 注释裁切，若干版式/图注不合格。"),
            ("提交快照可复现", "不通过", f"submission-v2 指向 HEAD，但工作树有 {AUDIT['git']['dirty_entry_count']} 个变更/未跟踪项；依赖无锁定，metadata 的 H3 版本与环境不一致。"),
            ("作者、贡献与 AI 声明", "待作者确认", "pyproject/README 残留参考论文机构；CRediT 未填；AI 声明与仓库中实际辅助痕迹可能不一致。"),
        ],
    )
    add_callout(
        doc,
        "推荐决定：Major Revision / 暂不投稿",
        "这不是因为所有结果都不可用，而是因为若按当前文字投稿，审稿人能从代码直接反证关键方法声明。修复顺序必须是：冻结真实数据来源→修复全量模型→统一 bbox/标签/尺度定义→全链重跑→再改文稿与图。",
        fill=PALE_ORANGE,
        accent=ORANGE,
    )

    add_heading(doc, "1.1 经过独立核验仍可保留的结果", 2)
    cv_lm = AUDIT["spatial_cv_independent_recompute"]["lower_manhattan"]
    cv_ex = AUDIT["spatial_cv_independent_recompute"]["expanded"]
    add_table(
        doc,
        ["数据集", "n", "Accuracy / F1", "R² / pooled ROC-AUC / AP"],
        [
            ("小范围（实际 smoke bbox）", "141", f"{cv_lm['spatial_cv_accuracy_mean']:.3f} ± {cv_lm['spatial_cv_accuracy_std']:.3f} / {cv_lm['spatial_cv_f1_mean']:.3f}", f"{cv_lm['spatial_cv_r2_mean']:.3f} ± {cv_lm['spatial_cv_r2_std']:.3f} / {cv_lm['spatial_cv_roc_auc_pooled']:.3f} / {cv_lm['spatial_cv_pr_auc_pooled']:.3f}"),
            ("扩展范围", "956", f"{cv_ex['spatial_cv_accuracy_mean']:.3f} ± {cv_ex['spatial_cv_accuracy_std']:.3f} / {cv_ex['spatial_cv_f1_mean']:.3f}", f"{cv_ex['spatial_cv_r2_mean']:.3f} ± {cv_ex['spatial_cv_r2_std']:.3f} / {cv_ex['spatial_cv_roc_auc_pooled']:.3f} / {cv_ex['spatial_cv_pr_auc_pooled']:.3f}"),
        ],
        [4.0, 1.0, 5.7, 6.3],
    )
    bullet(doc, "OOF 文件与训练表按 h3_index 一一对应，无重复网格，y_true 无错配。")
    bullet(doc, "目标构造 E=max(DEP 面积占比, 311 是否存在, HWM 是否存在) 在三个处理表中均可逐行复算。")
    bullet(doc, "测试结果：58 passed，1 skipped；src/scripts/tests compileall 通过。")
    bullet(doc, "当前 raw 快照关闭合成回填后可以重建处理表，说明现有主结果不是‘算一半、猜一半’。")

    add_heading(doc, "1.2 不能写成确定事实的内容", 2)
    bullet(doc, "不能写“所有原始数据是我们自己的”。原始输入是公开/第三方资料；作者自己的应是抓取快照、处理脚本、派生表、模型与图。")
    bullet(doc, "不能写“已证明绝无抄袭”。当前只对提供的参考论文做了局部文本与实现搜索；结果是‘未发现可疑实质复用’，不是绝对证明。")
    bullet(doc, "不能把模型输出称为物理洪水风险、积水深度或发生概率。当前监督目标是人为组合的开放证据分数。")
    bullet(doc, "不能把恒定 75 mm/h 的占位降雨写成已验证的 rainfall-conditioned model；四个场景的每个网格输出范围严格为 0。")

    section_break(doc)
    add_heading(doc, "2. 审查范围、方法与证据链", 1)
    text_para(doc, "审查覆盖论文 PDF/Markdown、报告 PDF/Markdown、src 与 scripts、tests、raw/processed 数据、模型工件、输出 CSV/JSON 与 8 组论文图。所有审查性脚本和重建输出都写入审查输出目录，没有覆盖原论文、原模型或原数据。")
    add_heading(doc, "2.1 执行的核验", 2)
    checks = [
        "静态代码审查：训练/保存路径、特征和标签组装、空间 CV、尺度诊断、情景预测、负对照、绘图及元数据。",
        "结构与哈希清点：文件数、raw 向量特征数/几何类型、处理表主键与空值、模型元数据、Git 快照。",
        "独立数值复算：从 OOF/折表重新计算 accuracy、F1、R²、ROC-AUC、AP，并与 metadata 对照。",
        "模型身份实验：分别用相同 80% split 和 100% 数据重拟合，与保存模型逐网格比较预测。",
        "raw→processed 重建：明确 fallback_synthetic=False，重新组装 141、956 与 R10 991 网格。",
        "图片重生成：调用当前绘图链生成审查副本，比较尺寸、哈希和数值来源；逐页检查 manuscript.pdf（23 页）与 report.pdf（29 页）。",
        "局部原创性检查：将论文正文与仓库提供的参考论文 Markdown 做 10-token 精确重叠和段落相似度检查，并搜索 PFIb 数据/模型痕迹。",
        "端点身份检查：查询 ArcGIS 服务 item owner/org/license；对官方落地页进行交叉确认。",
    ]
    for x in checks:
        bullet(doc, x)

    add_heading(doc, "2.2 证据等级", 2)
    add_table(
        doc,
        ["等级", "含义", "本报告示例"],
        [
            ("A：可重复实验", "由当前代码/数据重新执行或独立数值复算", "CV 指标、raw 重建、模型 80%/100% 身份对照"),
            ("B：文件与元数据证据", "由代码行、哈希、manifest、Git 状态直接支持", "bbox、端点 owner、dirty tree、缺许可/校验和"),
            ("C：人工视觉审查", "逐页/逐图观察，适合版式和语义判断", "Fig.4 裁切、表格跨页、公式未渲染"),
            ("D：有限排除证据", "未发现不等于绝对不存在", "未发现引用论文 PFIb 数据或实质文本复制"),
        ],
        [2.8, 6.2, 8.0],
    )
    add_callout(doc, "可证伪性原则", "凡是无法从当前仓库、数据快照或权威元数据证明的事项，报告统一标注为“待作者确认”或“尚不能证明”，不以猜测填补证据缺口。", fill=PALE_TEAL, accent=TEAL)

    section_break(doc)
    add_heading(doc, "3. 真实性、独立性与原创性审查", 1)
    add_heading(doc, "3.1 当前证据支持什么", 2)
    bullet(doc, "现有处理表可由当前 raw 文件重建，H3 主键、标签、类别、文本 provenance 字段一致；最大连续值差异来自几何浮点计算。")
    bullet(doc, "主空间 CV 指标可从 OOF 文件独立复算，结果与 metadata 完全一致。")
    bullet(doc, "仓库未发现参考论文 proprietary PFIb 数据表、其模型权重或实现；本项目使用的 DEP/311/HWM/地形/建筑/水系输入与参考论文的 proprietary building-level PFIb 路径不同。")
    bullet(doc, f"正文对比得到 {AUDIT['manuscript_reference_similarity']['exact_10_word_overlap_count']} 个 10-word 精确重叠，全部来自 Elsevier 常见 competing-interest 模板；最高段落相似度约 {AUDIT['manuscript_reference_similarity']['top_paragraph_matches'][0]['similarity']:.3f}，未发现实质段落直接复制。")

    add_heading(doc, "3.2 当前证据不能支持什么", 2)
    bullet(doc, "公开数据不是作者自有财产。正确表述是：作者独立下载/冻结公开快照，编写处理和建模代码并生成派生结果。")
    bullet(doc, "DEP/311 服务当前 owner 分别是“20221058@GT2022”和“jeichen_GIS”，无官方组织 ID、license 或 copyright 字段；只能称为 public ArcGIS mirrors，不能称为经验证的官方 NYC 端点。")
    bullet(doc, "局部查重不能排除对未纳入审查的所有论文、代码库或历史文件的重用；作者仍须提供贡献记录、实验日志和数据快照来源。")
    bullet(doc, "同一次工作树内生成的图、表、论文互相一致，只能证明内部一致性；raw 重建与独立模型实验才是更强证据。")

    add_heading(doc, "3.3 原始数据身份与风险", 2)
    endpoint_rows = []
    for key, label in (("dep_stormwater_mirror", "DEP stormwater"), ("streetfloodtime_311_mirror", "311 streetfloodtime"), ("building_view", "Building footprints")):
        e = SOURCES["arcgis_services"][key]
        verdict = "官方身份未证实" if key != "building_view" else "NYC 组织/条款可见"
        endpoint_rows.append((label, e.get("item_owner"), e.get("item_org_id") or "—", verdict))
    add_table(doc, ["图层", "ArcGIS owner", "org ID", "审查判定"], endpoint_rows, [3.6, 4.2, 4.0, 5.2])
    text_para(doc, "建筑端点由 maps.nyc.data 持有，位于 NYC ArcGIS 组织并带 NYC Terms；USGS Ida HWM 有官方数据发布页与 DOI/CC0 信息。DEP/311 应改从官方数据门户或官方落地页可追踪的下载链接获取，或在文中明确说明镜像身份并保存原始官方链路。", size=9.2)

    add_heading(doc, "3.4 必须建立的真实性证据包", 2)
    required = [
        "data_manifest.csv/json：每一层记录 official_landing_page、exact_download_url/service+layer、item owner/org、query、vintage、retrieved_utc、CRS、bbox、feature_count、size_bytes、sha256、license。",
        "run_manifest.json：Git commit、dirty=false、Python/OS、完整依赖锁、H3 版本、随机种子、配置文件 SHA、全部 raw/processed/model/output SHA。",
        "provenance_by_column.csv：每个 feature/label 列的 source file、算法、单位、聚合方式、缺失数、合成回填数（生产运行必须为 0）。",
        "execution_log.txt：从空 outputs/models 开始的命令、UTC 起止时间、退出码和测试结果；保留只读原始快照。",
        "figure_manifest.csv：图号→脚本→输入文件 SHA→输出图 SHA→图中展示的列/指标→论文引用位置。",
        "authorship_record.md：每名作者 CRediT、数据/代码/图表责任人、AI 工具实际用途及人工复核签字。",
    ]
    for x in required:
        bullet(doc, x)

    section_break(doc)
    add_heading(doc, "4. P0：投稿前必须修复的问题", 1)

    p0s = [
        (
            "P0-01 研究范围与 n=141 自相矛盾",
            "论文写 bbox≈[-74.02, 40.70, -73.97, 40.76]；该范围在 R9 产生 262 个网格。现有 n=141、地图坐标和 nyc_smoke 实际使用 [-74.02, 40.70, -73.98, 40.74]。",
            "这是研究区域、样本量和全部小范围结果的定义错误。读者按论文无法复现 141。",
            "二选一并全局统一：(A) 保留 141 结果，把区域明确改称 smoke pilot extent 并写真实 bbox；(B) 保留“Lower Manhattan”主张，用论文 bbox 重建 262 网格并重跑所有表图。推荐 B，若时间不足才选 A。",
        ),
        (
            "P0-02 ‘full-fit’模型实际只在 80% 数据上训练",
            "model.py:67–79 先 train_test_split，再只 fit(X_tr)；pipeline.py:71–107 保存该对象。保存模型与独立 80% refit 的预测最大差为 0，与 100% refit 最大差在 LM 为 0.819（分类）/0.736（回归），expanded 为 0.718/0.605。",
            "Fig.2(c)、方法 3.6/4.1、自适应加密及相关图注的‘all 141 cells’陈述可被代码直接反证。",
            "训练函数返回 evaluation artifacts 与 deployment artifacts 两套对象。CV/holdout 仅评价；评价完成后用全部 X,y 新建 classifier_full/regressor_full 并保存独立文件。图 2(c) 与 adaptive 只读 full 文件，metadata 写 fit_rows、training_h3_sha256、role。重跑全部依赖输出。",
        ),
        (
            "P0-03 DEP/311 被写成官方源，但当前端点身份不成立",
            "ArcGIS item owner/org/license 检查显示 DEP 与 311 分别来自个人/非机构账号，org ID 为空，许可和版权字段为空。现有 manifest 没有 exact URL、逐层检索时间、许可或 SHA。",
            "来源真实性、可再获取性和法律许可均不足；“direct NYC DEP/NYC 311 source”表述过强。",
            "从 NYC 官方开放数据门户/官方页面追溯正式数据集 ID；重新下载、校验并冻结。若只能使用镜像，文中明确‘public mirror, official identity not verified’，同时给官方语义对照和许可依据。",
        ),
        (
            "P0-04 生产数据组装默认允许合成回填",
            "assemble.py:192–268 先生成 synthetic/hash 特征，fallback_synthetic 默认 True，并可对观测 join 的 NaN 回填 synthetic；pipeline 多处默认也为 True。observed.add 在检查完整性前执行。",
            "未来任何缺失文件或 join 空值都可能悄悄变成合成数据，provenance 仍可能显示 observed。这正是‘做一半、臆断一半’的结构性风险。",
            "生产入口默认 fallback_synthetic=False；发现文件缺失、列缺失或任一 NaN 立即失败。合成模式只能由显式 --demo/--allow-synthetic 开启，输出 data_mode=synthetic，且论文流水线禁止读取。逐列记录 observed_count/synthetic_count/null_count。",
        ),
        (
            "P0-05 版本归档声明不真实",
            f"HEAD 与 submission-v2 指向 {AUDIT['git']['head'][:12]}，但当前工作树 dirty 且有 {AUDIT['git']['dirty_entry_count']} 个变更/未跟踪项；论文和结果属于未归档状态。依赖仅用 >=，无 lock；metadata 写 H3 4.5.0，而当前环境为 4.4.2。",
            "第三方无法从所谓 immutable tag 得到当前论文和数据结果；像素级图也无法复现。",
            "先整理并移除临时/无关文件；锁定依赖；从干净工作树全链执行；在执行后提交、打新 tag/DOI；manifest 写 commit 和所有工件 SHA。论文只引用新归档。",
        ),
        (
            "P0-06 作者机构、贡献与参考论文残留",
            "pyproject.toml 作者写 NTNU / 7Analytics / Ostfold University College，README 开头重复参考论文机构；论文 CRediT 仍是占位符。补充 ZIP 的 PII 与参考论文不一致且内容是无关的 VolturnUS_Model.prt。",
            "会造成作者身份、机构冒用、材料来源或论文独立性的严重疑问。",
            "由作者本人确认并填写真实姓名/单位/ORCID/CRediT；删除所有无权使用的机构残留；将无关 ZIP 隔离出投稿包并记录原因。未经权利人确认不得保留机构名。",
        ),
        (
            "P0-07 尺度热点的 tie-break 与‘matched budgets’不成立",
            "R10 有 991 网格，top 10%=99；阈值为 1，149 个网格并列最高，代码按 H3 字符串任取 99。投影到 R9/R8 后 fine hotspot 覆盖 56/17 个父网格，而 coarse set 只有 16/3，支持比为 3.5/5.67。",
            "Jaccard 对任意字符串 tie-break 敏感；不同分辨率各取 10% 网格并不等于相同面积/空间预算。",
            "用物理面积/投影后支持匹配预算；并列值采用 all-ties 加权、fractional membership 或重复随机 tie-break bootstrap，报告中位数与 95% CI。增加连续 rank correlation 或 area-weighted overlap。删除‘matched budgets’直到真正实现。",
        ),
        (
            "P0-08 Data/code availability 与 AI 使用声明不完整",
            "论文声称 raw→URL/license 映射和 immutable tag，实际 manifest/工作树不支持；AI 声明只写编辑，但仓库存在广泛的代码、审查和写作辅助痕迹。",
            "提交声明与可观察证据不一致，属于研究诚信风险。",
            "按期刊政策和实际使用如实重写。列明 AI 用于语言、代码 review/debug、测试设计、图表一致性与审查草稿（以作者真实情况为准）；说明所有代码实际执行、输出核验和最终责任均由作者承担。",
        ),
    ]
    for title, evidence, impact, fix in p0s:
        add_heading(doc, title, 2)
        text_para(doc, "证据：" + evidence, bold_prefix="证据：", after=3)
        text_para(doc, "影响：" + impact, bold_prefix="影响：", after=3)
        text_para(doc, "具体修改：" + fix, bold_prefix="具体修改：", after=6)

    section_break(doc)
    add_heading(doc, "5. 数据、标签与空间逻辑审查", 1)
    add_heading(doc, "5.1 raw→processed 重建结果", 2)
    lm = REBUILD["lower_manhattan"]
    ex = REBUILD["expanded"]
    r10 = REBUILD["r10_labels"]
    add_table(
        doc,
        ["对象", "H3/列/标签", "最大数值差", "判定"],
        [
            ("LM 141 R9", "集合、列、文本、class 全同", f"building {lm['numeric_max_abs_diff']['building_density']:.2e}; distance {lm['numeric_max_abs_diff']['dist_stream_m']:.2e}", "科学等价"),
            ("Expanded 956 R9", "集合、列、文本、class 全同", f"building {ex['numeric_max_abs_diff']['building_density']:.2e}; distance {ex['numeric_max_abs_diff']['dist_stream_m']:.2e}", "科学等价"),
            ("R10 labels 991", "集合、列、文本、class 全同", f"max area difference {max(r10['numeric_max_abs_diff'].values()):.2e}", "科学等价"),
        ],
        [3.3, 4.6, 6.1, 3.0],
    )
    text_para(doc, "说明：若使用严格 10⁻¹² 阈值，部分几何列会被标记不同；差异上限约 10⁻⁸，来自 Shapely/几何运算顺序，不是标签或数据内容漂移。报告因此判定为科学等价而非字节完全相同。")

    add_heading(doc, "5.2 原始向量快照", 2)
    vector_map = {Path(x["path"]).name + " @ " + Path(x["path"]).parent.name: x for x in SOURCES["raw_vectors"]}
    wanted = []
    for key, label in [
        ("dep_stormwater_flood.geojson @ nyc", "DEP current (LM)"),
        ("flooding_311.geojson @ nyc", "311 (LM)"),
        ("flooding_311.geojson @ nyc_expanded", "311 (expanded)"),
        ("usgs_ida_hwm.geojson @ nyc", "USGS HWM"),
        ("building_footprints.geojson @ nyc", "Buildings (LM)"),
        ("building_footprints.geojson @ nyc_expanded", "Buildings (expanded)"),
    ]:
        x = vector_map.get(key)
        if x:
            wanted.append((label, x["n_features"], ", ".join(f"{k}:{v}" for k, v in x["geometry_types"].items()), x["sha256"][:16] + "…"))
    add_table(doc, ["图层", "features", "geometry", "SHA-256（前16位）"], wanted, [4.2, 2.0, 4.3, 6.5])
    text_para(doc, "DEP 与 HWM 文件在两个 study 目录中哈希完全相同并不自动表示错误：它们是全域/复合几何，后续在 H3 overlay 时裁剪。但 expanded manifest 中仍出现“Lower Manhattan bbox subset”式说明，必须纠正。")

    add_heading(doc, "5.3 标签定义的科学含义", 2)
    add_code_block(doc, "a_DEP,c = area(DEP polygon ∩ cell c) / area(cell c)\nz_311,c = 1 if at least one complaint point lies in c, else 0\nz_HWM,c = 1 if at least one Ida high-water mark lies in c, else 0\nE_c = max(a_DEP,c, z_311,c, z_HWM,c)\nY_c = 1[E_c ≥ 10^-9]")
    bullet(doc, "代码与表格严格遵循该 max 规则；这是可保留的内部一致性。")
    bullet(doc, "一个 311 点或一个 HWM 点会把 E 直接置为 1，信息量与大面积 DEP 重叠不在同一量纲；E 的双峰和 149 个满分 R10 网格正由此产生。")
    bullet(doc, "311 为 2010–2014 的历史报告，DEP 为模型输出面，HWM 为 Ida 2021 事件点；时间和机制混合，E 不是单一事件真值。")
    bullet(doc, "建议变量改名 flood_evidence_score / flood_evidence_class；全文避免称 risk、severity、depth、probability 或 ground truth。")
    bullet(doc, "做敏感性：point presence 改成经时间/质量加权的点密度，DEP 与点源分别标准化；报告 max、weighted sum、source-specific multi-task 三种定义。")

    add_heading(doc, "5.4 面积和距离计算的坐标系", 2)
    text_para(doc, "features.py:165–170 与 labels.py:367–383 直接用 EPSG:4326 经纬度多边形的 Shapely area 做面积比例。小区域内分子/分母部分抵消，且重建结果稳定，但这不是严谨的面积计算。")
    bullet(doc, "在相交前将网格与源几何投影到 EPSG:2263（NY StatePlane Long Island）或合适等面积 CRS。")
    bullet(doc, "记录 CRS、单位、无效几何修复数和 area conservation check；用旧/新实现做差异表。")
    bullet(doc, "dist_stream_m 实际包含岸线、潮汐河流和水体，应改名 distance_to_mapped_water_m，不要写成经典 stream proximity。")

    add_heading(doc, "5.5 降雨条件", 2)
    rs = AUDIT["rainfall_scenarios"]
    add_table(doc, ["场景", "mm/h", "n", "mean PFI_h"], [(x["scenario"], x["rainfall_mm_h"], x["rows"], f"{x['mean_PFI_h']:.9f}") for x in rs["scenarios"]], [4.5, 3.0, 3.0, 6.5])
    text_para(doc, f"四场景共 {rs['rows']} 行；每个网格的 max−min={rs['max_within_cell_range']:.1f}，非零响应网格={rs['nonzero_range_cells']}。训练 rainfall_mm_h 恒为 75，模型中该特征重要度为 0。")
    add_callout(doc, "结论", "当前只能把 rainfall_mm_h 描述为接口占位/未识别条件，不能声称场景响应或事件条件化已被数据验证。若论文主线不依赖降雨，建议从主模型与主图中移除；若保留，必须引入具有空间/事件变化的真实雷达或雨量站数据并重新训练。", fill=PALE_YELLOW, accent=ORANGE)

    section_break(doc)
    add_heading(doc, "6. 代码与模型完整性审查", 1)
    add_heading(doc, "6.1 已做得较好的部分", 2)
    bullet(doc, "以 H3 R7 parent 作为 GroupKFold group，定义上避免同一 parent block 同时进入训练和测试。")
    bullet(doc, "OOF 预测逐网格保存，ROC-AUC/AP 可以在折外预测上重新计算；预处理位于折内 pipeline。")
    bullet(doc, "当前 paper run 显式给 negative control 传入 OOF model score，避免直接比较由定义决定的 target score。")
    bullet(doc, "原始重建在 fallback_synthetic=False 下无空值，说明当前结果所依赖的数据列是完整的。")

    add_heading(doc, "6.2 需要修改的代码问题", 2)
    code_rows = [
        ("严重", "model.py:67–101; pipeline.py:71–107", "保存的是 80% split 模型，不是 full fit", "分离 eval/full artifacts；metadata 写 fit role/rows/hash"),
        ("严重", "assemble.py:192–268; pipeline defaults", "生产默认 synthetic fallback，可能伪装为 observed", "默认 fail-closed；显式 demo 开关；逐列计数"),
        ("高", "features.py:165–170; labels.py:367–383", "经纬度平面 area", "投影到本地/等面积 CRS 后求交"),
        ("高", "scale diagnostic scripts", "149 个满分并列，字符串 tie-break；预算未按面积匹配", "weighted/all-tie/bootstrap；area-matched support"),
        ("中", "estimators.py/config.py", "estimator 使用模块常量 42，不读运行配置 seed", "seed 作为函数参数贯穿并写 metadata"),
        ("中", "baselines / negative_control helpers", "宽泛 except 与危险默认列可能隐藏错误", "只捕获预期异常；要求显式 score_col"),
        ("中", "metadata.py / requirements", "绝对路径、版本漂移、无锁文件", "相对路径、commit/hash、uv/pip lock"),
        ("中", "tests", "缺 raw→paper、full-fit、manifest、tie、视觉回归测试", "加入集成和不变量测试"),
    ]
    add_table(doc, ["级别", "位置", "问题", "修改"], code_rows, [1.6, 4.5, 6.1, 4.8], font_size=7.8)

    add_heading(doc, "6.3 推荐的模型工件协议", 2)
    add_code_block(doc, "models/<run>/\n  evaluation/\n    spatial_cv_folds.csv\n    spatial_cv_oof_predictions.csv\n    split_diagnostic_classifier.joblib   # optional; never used for maps\n  deployment/\n    classifier_full.joblib               # fit_rows == n_cells\n    regressor_full.joblib\n    feature_columns.json\n  run_manifest.json                      # role, fit_rows, H3 hash, data/config/code hashes")
    bullet(doc, "每个加载函数必须接受 expected_role；绘图/自适应要求 role=deployment_full，CV 汇总只读 evaluation。")
    bullet(doc, "新增测试：保存模型在全部 X 上的 prediction 必须等于同配置独立 100% refit，且 manifest.fit_rows==n_cells。")
    bullet(doc, "新增测试：production run 中 synthetic_count 或 null_count 非零立即失败。")

    add_heading(doc, "6.4 评价设计局限", 2)
    bullet(doc, "空间分块比随机拆分更诚实，但相邻 R7 block 仍共享边界；未设置 spatial buffer，不能声称完全无空间泄漏。")
    bullet(doc, "小范围仅 7 个 block/141 cells，折大小 21–49，方差大；扩展范围 28 blocks 更可信，但仍无外部城市/事件验证。")
    bullet(doc, "k=3/R6 只有 3 groups，因此最多 3 folds；Table 8 caption 的“same five-fold protocol”错误。")
    bullet(doc, "source ablation 只能说降低单一来源驱动的担忧，不能证明不存在 311 报告偏差、空间混杂或 DEP 预测面造成的循环。")

    section_break(doc)
    add_heading(doc, "7. 结果与模型的独立数值核验", 1)
    add_heading(doc, "7.1 保存模型身份实验", 2)
    sm = AUDIT["saved_model_identity"]
    rows = []
    for key, label in (("lower_manhattan", "LM"), ("expanded", "Expanded")):
        x = sm[key]
        rows.append((label, f"{x['independent_train_rows_80']}/{x['independent_train_rows_100']}", f"{x['stored_vs_refit80_max_abs_probability_diff']:.1f}", f"{x['stored_vs_refit100_max_abs_probability_diff']:.3f}", f"{x['stored_vs_refit100_max_abs_regression_diff']:.3f}"))
    add_table(doc, ["数据", "80%/100% rows", "saved vs 80% proba", "saved vs 100% proba", "saved vs 100% reg"], rows, [2.1, 3.2, 3.8, 4.0, 3.9], font_size=8.0)
    text_para(doc, "saved vs 80% 差为 0 是决定性证据：当前保存模型就是 split 后的训练子集模型。saved vs 100% 的大差异说明问题不是数值噪声或序列化差异。")

    add_heading(doc, "7.2 空间 CV 可复算性", 2)
    add_table(
        doc,
        ["检查", "LM", "Expanded", "判定"],
        [
            ("OOF 行数/重复", "141 / 0", "956 / 0", "通过"),
            ("y_true 与处理表错配", "0", "0", "通过"),
            ("pooled ROC-AUC", f"{cv_lm['spatial_cv_roc_auc_pooled']:.9f}", f"{cv_ex['spatial_cv_roc_auc_pooled']:.9f}", "与 metadata 一致"),
            ("pooled AP", f"{cv_lm['spatial_cv_pr_auc_pooled']:.9f}", f"{cv_ex['spatial_cv_pr_auc_pooled']:.9f}", "与 metadata 一致"),
            ("fold accuracy mean", f"{cv_lm['spatial_cv_accuracy_mean']:.9f}", f"{cv_ex['spatial_cv_accuracy_mean']:.9f}", "与 metadata 一致"),
            ("fold R² mean", f"{cv_lm['spatial_cv_r2_mean']:.9f}", f"{cv_ex['spatial_cv_r2_mean']:.9f}", "与 metadata 一致"),
        ],
        [4.1, 4.1, 4.1, 4.7],
    )
    text_para(doc, "注：metadata 的 pr_auc 实际是 sklearn average_precision_score，论文写 AP 是正确的；代码字段建议改名 pooled_average_precision，避免与梯形积分 PR-AUC 混淆。")

    add_heading(doc, "7.3 尺度诊断的再解释", 2)
    sc = AUDIT["scale_diagnostic"]
    add_table(
        doc,
        ["量", "值", "含义"],
        [
            ("R10 total / top-k", f"991 / {sc['r10_k']}", "名义 10%"),
            ("top-k threshold", f"{sc['r10_threshold']:.1f}", "最高分即阈值"),
            ("tied at threshold", str(sc["r10_ties_at_threshold"]), "149 个候选中任取 99"),
            ("projected fine parents R9 / coarse set", "56 / 16", "空间支持比 3.5"),
            ("projected fine parents R8 / coarse set", "17 / 3", "空间支持比 5.67"),
        ],
        [4.5, 4.2, 8.3],
    )
    text_para(doc, "所以当前 Jaccard 数值虽然可由代码得到，但其科学解释不是稳定的‘尺度损失’，而是特定 tie-break 和不等支持预算下的集合重叠。表和图应在方法修复后重算。")

    section_break(doc)
    add_heading(doc, "8. 图片—数据—文字的逐图审查", 1)
    add_heading(doc, "8.1 图表对应性总表", 2)
    fig_rows = [
        ("Fig.1 workflow", "当前代码流程基本对应", "PASS/改字", "把 rainfall-conditioned 改为 rainfall hook；明确 full deployment model 与 CV 分离"),
        ("Fig.2 spatial maps", "a=target；b=OOF；c=保存模型", "FAIL", "c 不是 all-cell full-fit；修复模型后重画并更新统计/图注"),
        ("Fig.3 source maps", "DEP/311/HWM/max 规则对应", "PASS/来源风险", "HWM 在 LM 为 0 与数据一致；端点身份及 target 命名需改"),
        ("Fig.4 spatial CV", "折分数据对应", "VISUAL FAIL", "底部 n/prevalence 注释被裁切/与图例冲突；增大 bottom margin 或移入面板"),
        ("Fig.5 multi-resolution", "R10→R9/R8 mean 表面对应", "PASS/解释风险", "图本身可重生；不要由其推出稳健 hotspot budget"),
        ("Fig.6 resolution effects", "CSV 数值对应", "FAIL/方法", "tie 与不等空间支持使 Jaccard 解释失效，方法修复后重画"),
        ("Fig.S1 Jaccard", "CSV 数值对应", "FAIL/方法", "同上；增加 tie bootstrap CI/area-weighted 指标"),
        ("Fig.S2 adaptive", "当前保存模型与 cell count 对应", "FAIL/语义", "使用错误的 80% 模型筛选；只证明表示规模，不证明精度/效率"),
    ]
    add_table(doc, ["图", "数据逻辑", "状态", "修改要求"], fig_rows, [3.0, 4.6, 2.2, 7.2], font_size=7.8)

    add_heading(doc, "8.2 Figure 2：视觉正确但模型语义错误", 2)
    add_figure(doc, ROOT / "docs" / "paper" / "figures" / "spatial_maps.png", "现有 Figure 2。面板 (a)/(b) 的数据角色可追踪；面板 (c) 的视觉表面来自保存的 80% split 模型，却在图注中写成 all-141 full fit。", 16.8)
    bullet(doc, "颜色范围 0–1、三面板空间范围和 H3 支持前后对应；视觉上没有明显错位。")
    bullet(doc, "真正问题是“漂亮且合理的图”掩盖了工件角色错误：视觉可信不能替代模型 provenance。")
    bullet(doc, "修复后必须更新：图像、mean/median、与 OOF 的 Pearson r、adaptive selected cells、正文 4.1/4.4、caption。")

    add_heading(doc, "8.3 Figure 4：可见裁切缺陷", 2)
    add_figure(doc, ROOT / "docs" / "paper" / "figures" / "spatial_cv_folds.png", "现有 Figure 4。每折 n/正类占比注释在坐标轴下缘被裁切，并与全局图例/页面边界争抢空间。该缺陷在重新生成的审查副本中仍存在。", 15.8)
    bullet(doc, "修图建议：使用 constrained_layout=True 或 fig.subplots_adjust(bottom≥0.16)；把 n/prevalence 放入 axes transform 的左上角小框；保存前执行 bbox_inches='tight' 并做像素边界测试。")
    bullet(doc, "新增视觉回归：PNG 四边非背景内容不得在 3–5 px 内触边；OCR/模板检查所有预期注释出现。")

    add_heading(doc, "8.4 像素复现与数值复现", 2)
    same_shape = sum(1 for x in FIG_REPRO["figure_comparisons"] if x["same_shape"])
    exact = sum(1 for x in FIG_REPRO["figure_comparisons"] if x["same_sha256"])
    text_para(doc, f"8 组图均从当前工作流生成了审查副本；其中 {same_shape} 组尺寸完全相同，{exact} 组 SHA-256 完全相同。其余图的宽高有 1–10 px 差异或字体栅格差异。用于图 6/S1 的 Jaccard CSV 数值匹配。")
    bullet(doc, "可声称 numerical reproducibility，但不能声称 pixel-identical reproducibility。")
    bullet(doc, "锁定 matplotlib、SciencePlots、字体文件、操作系统/渲染后端和 savefig 参数，才可能稳定像素哈希。")

    section_break(doc)
    add_heading(doc, "9. 论文与报告逐项审稿意见", 1)
    add_heading(doc, "9.1 论文结构与论述", 2)
    paper_rows = [
        ("标题/摘要", "高", "risk/susceptibility 可能超过 evidence target 的含义；摘要应突出 open-evidence screening 和无外部验证。"),
        ("2.1 Study area", "严重", "bbox 与 n=141 冲突；按 P0-01 二选一。"),
        ("2.4 Target", "高", "解释 max 规则、点存在即 1、时间混合和非物理风险。"),
        ("3.3 Modelling", "严重", "写清 CV model、random-split diagnostic、full deployment model 三者完全分离。"),
        ("3.5 Scale", "严重", "删除 matched budget；增加 tie/area matching 方法。"),
        ("3.6 Adaptive", "严重", "改用全量模型；仍只能称 representation-size experiment。"),
        ("4.1 Fig.2", "严重", "修复 full-fit 后重写所有描述性统计。"),
        ("4.5 Rainfall", "高", "明确为 invariant placeholder；不以 PFI_h(c,r) 作为已验证贡献。"),
        ("4.8 Ablation", "中", "将‘proves not artifact’改为‘reduces concern’; 承认 reporting bias/confounding。"),
        ("Table 8 caption", "中", "改为 up to five folds；R6/k=3 仅三折。"),
        ("Conclusion", "中", "删除‘block-size sensitivity remains future’，因为 4.9 已完成。"),
        ("Availability/AI/CRediT", "严重", "按真实归档、实际工具使用和实际作者信息填写。"),
    ]
    add_table(doc, ["位置", "级别", "意见"], paper_rows, [4.1, 2.0, 10.9], font_size=8.1)

    add_heading(doc, "9.2 PDF 视觉与排版", 2)
    bullet(doc, "manuscript.pdf 23 页已逐页检查：第 9 页 Fig.2 独占上方且下部大面积空白；第 11 页 Fig.4 注释裁切；Table 7 跨第 16–17 页但无重复表头/continued；第 20 页 CRediT 占位符明显。")
    bullet(doc, "report.pdf 29 页已逐页检查：第 6 页公式保留原始 \\[ 标记并未渲染；第 10 页大面积留白；第 12 页 Fig.4 同样裁切；末页近乎空白。")
    bullet(doc, "解决办法：不要靠手工 page-break 固定图；用浮动/保持段落联动，长表启用 repeat header，公式用 MathML/Word equation 或生成 SVG，最终逐页渲染 QA。")

    add_heading(doc, "9.3 README、临时文件与提交包", 2)
    bullet(doc, "README 摘要保留旧指标（如 expanded accuracy 0.642、R² 0.525、ROC-AUC 0.70、adaptive 0.569），与当前论文 0.821/0.333/0.883/0.701 不符。")
    bullet(doc, "docs/paper/audit.md 包含多轮旧结果，不能作为当前独立证据；应按版本归档而不是持续覆盖混写。")
    bullet(doc, "scripts/_tmp_*.py、_redownload_dep.py 和无关 supplement ZIP 不应进入投稿包。先隔离，不要未经确认直接删除用户文件。")

    section_break(doc)
    add_heading(doc, "10. 具体修改路线与验收顺序", 1)
    phases = [
        ("阶段 1｜冻结证据", "确定真实作者/机构；从官方入口重取或明确镜像；生成含 URL/许可/SHA 的 manifest；选择真实 bbox；创建干净分支和依赖锁。"),
        ("阶段 2｜修复计算", "生产 fail-closed；投影面积；分离 full model；修复 seed/metadata；实现 tie-aware、area-matched scale diagnostics。"),
        ("阶段 3｜全链重跑", "从只读 raw 快照重建 processed→CV→full model→scenario/adaptive/scale→tables→figures；保存命令和退出码。"),
        ("阶段 4｜自动一致性", "加 manuscript-number check、figure manifest、full-fit identity、raw checksum、no-synthetic、visual boundary tests。"),
        ("阶段 5｜重写论文", "先方法和数据，再结果，再摘要/结论/图注；所有数字只从单一 machine-readable results.json 注入。"),
        ("阶段 6｜提交冻结", "测试全过；工作树 clean；提交并打新 tag/DOI；再次从归档 clone 重跑 smoke+key metrics；最后逐页 PDF QA。"),
    ]
    add_table(doc, ["阶段", "完成标准"], phases, [4.2, 12.8], font_size=8.5)

    add_heading(doc, "10.1 建议新增的自动验收门", 2)
    gates = [
        "assert production_manifest.synthetic_value_count == 0",
        "assert all(raw_layer.sha256 and raw_layer.exact_url and raw_layer.license)",
        "assert deployment_model.fit_rows == processed_table.n_rows",
        "assert deployment_model.training_h3_sha256 == sha256(sorted(h3_index))",
        "assert manuscript_bbox_cell_count == processed_table.n_rows",
        "assert oof.h3_set == processed.h3_set and oof.y_true == processed.flood_class",
        "assert recomputed_metrics == paper_results_json within declared tolerance",
        "assert hotspot_budget is area/support matched and tie sensitivity CI exists",
        "assert figure edges pass clipping detector and every figure has manifest row",
        "assert git_dirty is false and git_commit == archived_release_commit",
    ]
    add_code_block(doc, "\n".join(gates))

    add_heading(doc, "10.2 修复后最少需要重跑的产物", 2)
    for x in [
        "data/processed/nyc_lower_manhattan_h3.*（或明确改名 smoke）与 nyc_expanded_h3.*",
        "models/*/evaluation OOF/folds + deployment full models + run_manifest",
        "outputs/pfi_h_scenarios、adaptive_vs_fixed_ablation、jaccard/resolution/hotspot sensitivity、block/source ablations",
        "Fig.2、Fig.4、Fig.6、Fig.S1、Fig.S2；依赖 full model 的所有描述性统计",
        "manuscript.md/html/pdf、report.md/html/pdf、README、Data/code availability、AI declaration、CRediT",
    ]:
        bullet(doc, x)

    section_break(doc)
    add_heading(doc, "11. 可直接替换的英文方法与声明文本", 1)
    add_callout(doc, "使用前提", "以下文字按当前论文语气撰写，但包含条件分支。只有完成相应代码/数据修复后才可粘贴；不能用文字掩盖未修复的实现。", fill=PALE_YELLOW, accent=ORANGE)

    add_heading(doc, "11.1 Study area（若保留当前 141-cell 结果）", 2)
    add_code_block(doc, "The smaller pilot used the configured smoke-test extent (74.02–73.98°W, 40.70–40.74°N), which yielded 141 H3 resolution-9 cells. The expanded Manhattan extent covered 74.03–73.94°W and 40.68–40.80°N and yielded 956 resolution-9 cells. The smaller extent is described as a computational pilot rather than as a complete representation of Lower Manhattan. All study-area coordinates and cell counts were generated from the same configuration files used by the analysis.")
    text_para(doc, "若坚持使用论文原 bbox，则不能使用这段，必须用 262-cell 重跑后填写新 n 和新结果。", color=RED, size=8.8)

    add_heading(doc, "11.2 Evidence target", 2)
    add_code_block(doc, "For each H3 cell c, we computed the fraction of the cell intersecting the current-sea-level DEP stormwater-flood polygon (a_DEP,c) and binary indicators for the presence of at least one 311 street-flooding report (z_311,c) and at least one USGS Ida high-water mark (z_HWM,c). The composite evidence score was E_c = max(a_DEP,c, z_311,c, z_HWM,c), and the binary evidence class was Y_c = 1(E_c > 0). E_c is an open-evidence screening target: it combines model-derived polygons and event/report presence from different periods and is not interpreted as inundation depth, physical severity, annual probability, or ground truth. Source-specific columns were retained to expose the contribution of each component.")

    add_heading(doc, "11.3 Spatial evaluation and final model（仅在代码修复后使用）", 2)
    add_code_block(doc, "Predictive performance was estimated exclusively from out-of-fold predictions. Resolution-9 cells were grouped by their resolution-7 parents, and GroupKFold withheld complete parent blocks. All fitted transformations and estimators were trained within each training fold. ROC-AUC and average precision were computed from pooled out-of-fold scores; accuracy, F1, and evidence-score R² were summarized across folds. After evaluation was complete, separate deployment classifiers and regressors were refitted on all cells. These all-cell models were used only for descriptive maps and adaptive-grid screening and were never used to calculate held-out performance metrics. Model manifests record the model role, number and hash of training cells, configuration hash, software versions, and random seed.")

    add_heading(doc, "11.4 Provenance and fail-closed assembly（仅在 manifest/代码修复后使用）", 2)
    add_code_block(doc, "Each raw snapshot was accompanied by a machine-readable provenance record containing the official landing page, exact endpoint and layer or download URL, query parameters, data vintage, UTC retrieval time, coordinate reference system, requested extent, feature count, file size, SHA-256 checksum, and license. Production assembly was fail-closed: missing files, missing columns, failed joins, and non-finite values stopped the run. Synthetic fallback was disabled for all reported analyses and was available only through an explicitly labeled demonstration mode. Per-column provenance and missing-value summaries were written with each processed table.")

    add_heading(doc, "11.5 Scale-loss analysis（仅在 tie/预算修复后使用）", 2)
    add_code_block(doc, "Hotspot comparisons used matched physical support after projecting fine-resolution cells to each coarser resolution. Because the fine evidence surface contained ties at the hotspot threshold, we did not select cells by lexicographic H3 order. Tied cells were handled using [fractional membership / all-tie weighting / repeated random tie breaking], and overlap statistics are reported as [weighted estimates / bootstrap medians with 95% intervals]. We additionally report area-weighted overlap and rank-based agreement so that conclusions do not depend on a single arbitrary top-k realization.")

    add_heading(doc, "11.6 Rainfall limitation", 2)
    add_code_block(doc, "Rainfall was represented by a spatially constant 75 mm h−1 placeholder in the present pilot. Consequently, rainfall variation was not identifiable during training, the fitted rainfall importance was zero, and predictions were invariant across the evaluated rainfall values. The scenario interface is therefore a software hook rather than an empirically validated rainfall-response model; no rainfall-conditioning claim is made until spatially and temporally varying observed rainfall is introduced.")

    add_heading(doc, "11.7 Corrected Figure 2 caption（修复 full model 后）", 2)
    add_code_block(doc, "Figure 2. Spatial results for the [141-cell smoke pilot / corrected Lower Manhattan analysis]. (a) Composite open-evidence score; (b) pooled out-of-fold gradient-boosting score under H3-block spatial cross-validation; and (c) descriptive score from a separate model refitted on all analysis cells after evaluation. Panel (c) was not used to estimate predictive performance. The rainfall input is constant in the present pilot, so the mapped surface is invariant across the evaluated rainfall values.")

    add_heading(doc, "11.8 Corrected Table 8 caption", 2)
    add_code_block(doc, "Table 8. Sensitivity to the H3 parent resolution used for spatial blocking. Grouped cross-validation used up to five folds, with the number of folds limited by the number of available parent groups; the R6 configuration with three groups therefore used three folds.")

    add_heading(doc, "11.9 Honest data/code availability（提交归档前的工作稿）", 2)
    add_code_block(doc, "The analysis snapshot audited for this manuscript was a working tree and was not yet an immutable release. Before publication, the authors will deposit a clean, versioned archive containing the analysis code, configuration files, dependency lock, machine-readable provenance manifests, and checksums for all redistributable inputs and derived artifacts. Where licenses prevent redistribution, the archive will provide the official landing page, exact retrieval query, vintage, checksum, and reconstruction instructions. The final manuscript will cite the archived commit and DOI generated after the complete workflow has been rerun from the frozen inputs.")

    add_heading(doc, "11.10 Generative-AI declaration（必须按实际情况由作者确认）", 2)
    add_code_block(doc, "During the preparation of this work, the authors used ChatGPT for language editing and, where applicable, code review and debugging, test design, figure–result consistency checks, and drafting audit suggestions. The tool was not treated as a source of scientific evidence and did not independently execute or validate the reported experiments. The authors executed the code, inspected the source data, verified the numerical outputs, revised all generated text, and take full responsibility for the content of the manuscript. [Revise this statement to match the journal policy and the authors’ complete, factual usage record.]")

    section_break(doc)
    add_heading(doc, "12. 修复完成后的真实性验收清单", 1)
    checklist = [
        ("□", "来源", "每层有官方 landing page、exact endpoint/query、owner/org、license、retrieved_utc、SHA-256；镜像明确标注。"),
        ("□", "数据", "raw 只读；processed 从 raw 重建；主键/行数/空值/单位/CRS 通过；synthetic_count=0。"),
        ("□", "范围", "配置 bbox、论文坐标、地图坐标和 H3 cell count 完全一致。"),
        ("□", "标签", "max/替代 target 规则逐行可复算；变量命名不夸大；时间/来源局限写清。"),
        ("□", "模型", "OOF 仅用于评价；deployment model 确认 fit_rows=n；两类工件不混用。"),
        ("□", "尺度", "处理阈值并列；预算按面积/投影支持匹配；提供敏感性区间。"),
        ("□", "指标", "所有论文数字由 results.json 自动注入，并由独立脚本复算。"),
        ("□", "图", "每图有 input/output hash；图注角色正确；无裁切、重叠、误导色标。"),
        ("□", "版本", "依赖锁定、Git clean、commit/tag/DOI 一致；从归档克隆可复跑关键结果。"),
        ("□", "作者", "真实 author/affiliation/ORCID/CRediT；无参考论文机构残留；AI 声明完整。"),
        ("□", "投稿包", "无临时脚本、旧指标 README、混杂 audit 历史或无关 supplement。"),
        ("□", "人工签字", "数据、代码、统计、图片、写作各责任人完成最终逐项核对。"),
    ]
    add_table(doc, ["完成", "维度", "验收标准"], checklist, [1.4, 2.6, 13.0], font_size=8.4)
    add_callout(doc, "最终放行条件", "全部 P0 关闭、所有自动门通过、工作树干净、从冻结快照重新生成的论文 PDF 逐页检查无缺陷后，才建议将审稿结论从“重大修改”改为“可投稿”。", fill=PALE_RED, accent=RED)

    section_break(doc)
    add_heading(doc, "13. 审查证据索引", 1)
    evidence_files = [
        (EVIDENCE / "audit_evidence.json", "Git/metadata/processed/OOF/model identity/rainfall/scale/originality综合证据"),
        (EVIDENCE / "raw_rebuild_evidence.json", "关闭 synthetic fallback 的 raw→processed 重建对照"),
        (EVIDENCE / "source_endpoint_audit.json", "ArcGIS item owner/org/license 与 raw 向量快照"),
        (EVIDENCE / "file_sha256_inventory.csv", "项目文件 SHA-256 清单"),
        (OUT_DIR / "figure_reproduction" / "figure_reproduction_evidence.json", "8 组图片重生成与哈希/尺寸对照"),
        (OUT_DIR / "audit_evidence.py", "综合审查脚本"),
        (OUT_DIR / "rebuild_from_raw.py", "raw 重建脚本"),
        (OUT_DIR / "profile_sources.py", "端点与 raw 快照审查脚本"),
        (OUT_DIR / "reproduce_figures.py", "图表重生成脚本"),
    ]
    rows = []
    for p, purpose in evidence_files:
        # The section heading already fixes the root at 审查输出; basenames keep
        # the audit index readable and prevent an otherwise nearly blank last page.
        rows.append((p.name, p.stat().st_size if p.exists() else "MISSING", sha256(p)[:20] + "…" if p.exists() else "—", purpose))
    add_table(doc, ["路径", "bytes", "SHA-256（前20位）", "用途"], rows, [6.0, 2.0, 4.4, 4.6], font_size=7.4)

    add_heading(doc, "13.1 关键外部落地页（供修复 provenance 使用）", 2)
    links = [
        ("NYC AdaptNYC / official stormwater maps", "https://www.nyc.gov/content/climate/pages/initiatives/adaptnyc"),
        ("USGS Ida high-water marks data release", "https://www.usgs.gov/data/high-water-marks-five-boroughs-new-york-city-flash-flooding-caused-remnants-hurricane-ida"),
        ("NYC Open Data Building Footprints", "https://data.cityofnewyork.us/City-Government/Building-Footprints/3g6p-4u5s"),
        ("NYC Building Footprints metadata", "https://github.com/CityOfNewYork/nyc-geo-metadata/blob/main/Metadata/Metadata_BuildingFootprints.md"),
        ("Supplied reference paper landing page", "https://papers.ssrn.com/sol3/papers.cfm?abstract_id=5875380"),
    ]
    add_table(doc, ["资源", "URL"], links, [6.0, 11.0], font_size=7.7)

    add_heading(doc, "13.2 结论的适用边界", 2)
    bullet(doc, "本报告证明的是当前本地快照可在现有环境中重建/复算，不代表上游数据本身没有测量偏差、模型误差或许可问题。")
    bullet(doc, "未发现参考论文数据/文本的实质直接复用，不等于对所有外部来源完成法证级查重。")
    bullet(doc, "作者身份、真实贡献、数据获取授权与 AI 使用记录必须由作者提供并签字确认，代码审查不能替代这一责任；修复任何 P0 计算逻辑后，旧图、旧表和本报告中的数值都不能直接沿用，必须由新冻结快照重新生成。")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("— END OF AUDIT —")
    set_run_font(r, name="Arial", size=8, bold=True, color=TEAL)

    doc.core_properties.title = "严格审稿与真实性证据报告"
    doc.core_properties.subject = "Pluvial flood risk DGGS-H3 code/data/paper audit"
    doc.core_properties.author = "Independent technical audit generated in Codex; author verification required"
    doc.core_properties.keywords = "audit, reproducibility, provenance, H3, flood evidence"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
